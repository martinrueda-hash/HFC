"""Build the Excel data template for the HFC Market Intelligence Dashboard.

Usage:  python3 tools/build_excel_template.py [oxford_economics_export.csv]
Output: templates/HFC_market_intelligence_template.xlsx
Then:   recalculate (LibreOffice) and run tools/sync_data_js.py to refresh dashboard/data.js.

Sheets
  Read me           how to fill it in and upload it
  Settings          company, year, currency, Oxford Economics indicator names
  Oxford Economics  paste area for the Oxford Economics export (level values, as downloaded)
  Market levels     formulas: construction work done per country / end market / year
  Raw data          one row per country x end market; market growth is calculated from Market levels
  Dashboard data    formulas only: the table the dashboard reads on upload

End markets: Residential = Residential total - Multi family; Non-residential = Multi family + Commercial.
"""
import csv
import json
import math
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "templates" / "HFC_market_intelligence_template.xlsx"
OE_CSV = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "oxford-economics" / "Oxecon_download_-_9_October_2026.csv"

# Palette from the EU WT strategy deck
RED, RED_SOFT, SALMON, BLUSH, TINT = "CE0E2D", "DD6556", "E49587", "F4E4DC", "FCF1F1"
INK, INK_STRONG, MUTED, LINE, GREY_BAND = "505050", "191919", "8C8C8C", "E4E4E4", "7A7A7A"
INPUT_BLUE, LINK_GREEN, WHITE = "0000FF", "008000", "FFFFFF"
FONT = "Arial"
N_ROWS = 60          # rows available in Raw data
OE_LAST_ROW = 1000   # paste area depth on the Oxford Economics sheet
OE_LAST_COL = "AZ"

IND_TOTAL = "Work Done - Residential - Total"
IND_MF = "Work Done - Residential - Multi Family"
IND_COM = "Work Done - Non-Residential Building - Commercial"

# Countries: display name, code, Oxford Economics location, and ILLUSTRATIVE HFC inputs
# (market share 2025 res / non-res %, HFC CAGR hist / growth CY / BP CAGR %, scores prof / comp / channel / fit)
COUNTRIES = [
    ("Poland", "PL", "Poland", (5.0, 3.8), (9.0, 8.5, 10.0), (3, 3, 2, 4)),
    ("Czech Republic", "CZ", "Czech Republic", (6.0, 4.0), (6.0, 5.0, 7.0), (3, 3, 3, 3)),
    ("Hungary", "HU", "Hungary", (4.0, 3.0), (5.0, 3.0, 6.0), (2, 3, 2, 3)),
    ("Netherlands", "NL", "Netherlands", (16.0, 13.5), (3.5, 2.0, 3.0), (4, 3, 5, 4)),
    ("Sweden", "SE", "Sweden", (7.0, 5.5), (1.0, 0.0, 3.5), (4, 3, 3, 3)),
    ("Norway", "NO", "Norway", (7.5, 5.5), (1.5, 0.5, 3.0), (4, 3, 3, 3)),
    ("Finland", "FI", "Finland", (6.5, 5.0), (0.5, -0.5, 3.0), (4, 3, 3, 3)),
    ("Denmark", "DK", "Denmark", (8.0, 6.0), (2.0, 1.0, 3.0), (4, 3, 3, 3)),
    ("Germany", "DE", "Germany", (6.0, 4.5), (2.0, -2.0, 3.0), (3, 2, 3, 3)),
    ("Austria", "AT", "Austria", (7.0, 5.0), (2.0, -1.0, 2.5), (4, 3, 3, 4)),
    ("Switzerland", "CH", "Switzerland", (12.0, 10.0), (2.5, 1.5, 2.0), (5, 4, 4, 4)),
    ("United Kingdom", "UK", "United Kingdom", (3.2, 2.8), (-1.0, 1.5, 2.5), (3, 2, 2, 3)),
    ("France", "FR", "France", (18.0, 14.0), (4.0, 1.0, 3.5), (4, 3, 4, 4)),
    ("Italy", "IT", "Italy", (11.0, 9.5), (3.0, -3.0, 2.0), (3, 3, 4, 3)),
    ("Belgium", "BE", "Belgium", (15.0, 12.0), (3.0, 2.0, 3.5), (4, 3, 4, 4)),
    ("Spain", "ES", "Spain", (9.8, 7.5), (7.5, 6.0, 8.0), (3, 3, 3, 4)),
    ("Portugal", "PT", "Portugal", (8.0, 6.0), (6.5, 5.0, 5.0), (3, 3, 3, 4)),
]
# Illustrative addressable market: this share of construction work done (US$ m, 2023 prices) converted at 0.92 EUR/US$
ADDRESSABLE_RATIO = 0.012
USD_EUR = 0.92

def font(size=10, bold=False, color=INK_STRONG, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)

def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM_RED = Border(bottom=Side(style="medium", color=RED))
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

def band(ws, row, last_col, text, sub=None):
    """Red title band across the top of a sheet, like the deck's cover."""
    for c in range(1, last_col + 1):
        ws.cell(row=row, column=c).fill = fill(RED)
        ws.cell(row=row + 1, column=c).fill = fill(RED)
    ws.cell(row=row, column=1, value=text).font = Font(name=FONT, size=16, bold=True, color=WHITE)
    if sub:
        ws.cell(row=row + 1, column=1, value=sub).font = Font(name=FONT, size=10, color=WHITE)
    ws.row_dimensions[row].height = 28
    ws.row_dimensions[row + 1].height = 18

def section(ws, row, text, col=1):
    ws.cell(row=row, column=col, value=text).font = font(11, True, RED)

def input_cell(c, fmt=None):
    c.font = font(10, color=INPUT_BLUE)
    c.fill = fill(TINT)
    c.border = BOX
    if fmt:
        c.number_format = fmt

def calc_cell(c, fmt=None, size=10):
    c.font = font(size, color=LINK_GREEN)
    c.border = Border(bottom=Side(style="thin", color=LINE))
    if fmt:
        c.number_format = fmt

def year_formula(head, year_ref):
    """Turn 'CAGR {hy}-{py}' into an Excel string formula driven by the current year."""
    if "{" not in head:
        return head
    parts = re.split(r"(\{hy\}|\{py\}|\{cy\}|\{fy\})", head.replace("\n", " "))
    expr = {"{hy}": f"({year_ref}-6)", "{py}": f"({year_ref}-1)", "{cy}": year_ref, "{fy}": f"({year_ref}+5)"}
    return "=" + "&".join(expr[p] if p in expr else json.dumps(p) for p in parts if p != "")

# ---------------------------------------------------------------- Oxford Economics export
with open(OE_CSV, encoding="utf-8-sig", newline="") as fh:
    oe_rows = list(csv.reader(fh))
oe_head = oe_rows[0]

def to_cell(v, col_name):
    if re.fullmatch(r"\d{4}", col_name.strip()):
        try:
            return float(v)
        except ValueError:
            return None
    return v

def oe_level(location, indicator, year):
    yi = oe_head.index(str(year))
    for r in oe_rows[1:]:
        if r[0] == location and r[1] == indicator:
            return float(r[yi])
    return 0.0

wb = Workbook()

# ================================================================ Settings
st = wb.active
st.title = "Settings"
st.sheet_properties.tabColor = INK
band(st, 1, 4, "Settings", "Company, year and Oxford Economics series. Matrix weights and zone lines are set in the dashboard.")
st.column_dimensions["A"].width = 40
st.column_dimensions["B"].width = 46
st.column_dimensions["C"].width = 38
st.column_dimensions["D"].width = 60
for col, h in zip("ABCD", ["Setting", "Value", "Key (used by the dashboard, do not edit)", "Notes"]):
    c = st[f"{col}4"]
    c.value = h
    c.font = font(10, True, RED)
    c.border = BOTTOM_RED

SETTINGS = [
    ("General", None, None, None, None),
    ("Company name", "HFC", "company", "@", "Shown in column headers and labels"),
    ("Current year", 2026, "currentYear", "0", "Base year = current year - 1 (sizes and sales); history = 5 years to the base year; outlook = current year + 5"),
    ("Currency label", "EUR m", "currency", "@", "Unit of addressable market size and HFC sales"),
    ("Oxford Economics series (Indicator column of the export)", None, None, None, None),
    ("Residential - total", IND_TOTAL, "oe.res_total", "@", "Residential end market = Residential total - Multi family"),
    ("Residential - multi family", IND_MF, "oe.res_mf", "@", "Counted in Non-residential, as agreed"),
    ("Non-residential - commercial", IND_COM, "oe.nonres_com", "@", "Non-residential end market = Multi family + Commercial"),
]
r = 5
group_rows = []
SET_ROW = {}
for label, value, key, fmt, note in SETTINGS:
    if value is None:
        if r > 5:
            r += 1
        section(st, r, label)
        r += 1
        group_rows.append([])
        continue
    st.cell(row=r, column=1, value=label).font = font(10, label == "Total")
    if value == "SUM":
        rows = group_rows[-1]
        v = st.cell(row=r, column=2, value=f"=SUM(B{rows[0]}:B{rows[-1]})")
        v.font = font(10, True)
        v.number_format = fmt
        v.border = BOX
    else:
        v = st.cell(row=r, column=2, value=value)
        input_cell(v, fmt)
        group_rows[-1].append(r)
    if key:
        st.cell(row=r, column=3, value=key).font = font(8, color=MUTED)
        SET_ROW[key] = r
    st.cell(row=r, column=4, value=note).font = font(9, color=INK, italic=True)
    st.cell(row=r, column=4).alignment = WRAP
    r += 1
YEAR = f"Settings!$B${SET_ROW['currentYear']}"
IND = {k: f"Settings!$B${SET_ROW[k]}" for k in ("oe.res_total", "oe.res_mf", "oe.nonres_com")}
st.freeze_panes = "A5"

# ================================================================ Oxford Economics (paste area)
oe = wb.create_sheet("Oxford Economics")
oe.sheet_properties.tabColor = GREY_BAND
band(oe, 1, 16, "Oxford Economics", "Paste the full Oxford Economics download here, header row in row 6 (cell A6). Level values, any years.")
oe.cell(row=3, column=1, value=("To update: select A6, paste the new export (all columns, header row included) over the old one, and delete any leftover rows below. "
                                "Columns needed: Location, Indicator and one column per year. Nothing else to change.")).font = font(9, color=RED, italic=True)
oe.cell(row=4, column=1, value=f"Current data: {OE_CSV.name}").font = font(9, color=INK, italic=True)
OE_HEAD = 6
for j, h in enumerate(oe_head, 1):
    c = oe.cell(row=OE_HEAD, column=j, value=int(h) if re.fullmatch(r"\d{4}", h) else h)
    c.font = font(9, True, WHITE)
    c.fill = fill(INK)
    c.alignment = CENTER
for i, row in enumerate(oe_rows[1:], OE_HEAD + 1):
    for j, v in enumerate(row, 1):
        c = oe.cell(row=i, column=j, value=to_cell(v, oe_head[j - 1]))
        c.font = font(9, color=INPUT_BLUE)
        if isinstance(c.value, float):
            c.number_format = "#,##0.0"
oe.column_dimensions["A"].width = 18
oe.column_dimensions["B"].width = 46
for j in range(3, len(oe_head) + 1):
    oe.column_dimensions[get_column_letter(j)].width = 11
oe.freeze_panes = oe[f"C{OE_HEAD + 1}"]
OE = "'Oxford Economics'!"
OE_DATA = f"{OE}$A${OE_HEAD + 1}:${OE_LAST_COL}${OE_LAST_ROW}"
OE_HDR = f"{OE}$A${OE_HEAD}:${OE_LAST_COL}${OE_HEAD}"
OE_LOC = f"{OE}$A${OE_HEAD + 1}:$A${OE_LAST_ROW}"
OE_IND = f"{OE}$B${OE_HEAD + 1}:$B${OE_LAST_ROW}"

# ================================================================ Raw data
rd = wb.create_sheet("Raw data")
rd.sheet_properties.tabColor = RED
COLS = [
    # key, header, number format, group, width, comment, kind ("in" input / "calc" formula)
    ("country", "Country", "@", "Country", 17, None, "in"),
    ("code", "Code", "@", "Country", 7, "Short code (optional)", "in"),
    ("segment", "End market", "@", "Country", 15, "Residential, or Non-residential (multi family + commercial)", "in"),
    ("oe_location", "Oxford Economics\nlocation", "@", "Country", 16, "Location name exactly as in the Oxford Economics export", "in"),
    ("market_size", "Addressable market\n{py} (EUR m)", "#,##0", "Market", 13, "HFC addressable market in the base year (your estimate, not total construction output)", "in"),
    ("oe_level_py", "Construction work\ndone {py} (US$ m)", "#,##0", "Market", 14, "From Market levels (Oxford Economics, 2023 prices)", "calc"),
    ("mkt_cagr_hist", "Market CAGR\n{hy}-{py}", "0.0%", "Market", 12, "From Market levels: needs the export to include the start year", "calc"),
    ("mkt_growth_cy", "Market growth\n{cy}E", "0.0%", "Market", 12, "From Market levels", "calc"),
    ("mkt_cagr_fwd", "Market CAGR\n{cy}-{fy}", "0.0%", "Market", 12, "From Market levels", "calc"),
    ("reno_share", "Renovation share\nof market", "0%", "Split", 12, "Optional. Renovation as % of the market; new build is the rest", "in"),
    ("mkt_reno_cagr_hist", "Renovation CAGR\n{hy}-{py}", "0.0%", "Split", 12, "Optional", "in"),
    ("mkt_reno_growth_cy", "Renovation growth\n{cy}E", "0.0%", "Split", 12, "Optional", "in"),
    ("mkt_reno_cagr_fwd", "Renovation CAGR\n{cy}-{fy}", "0.0%", "Split", 12, "Optional", "in"),
    ("mkt_nb_cagr_hist", "New build CAGR\n{hy}-{py}", "0.0%", "Split", 12, "Optional", "in"),
    ("mkt_nb_growth_cy", "New build growth\n{cy}E", "0.0%", "Split", 12, "Optional", "in"),
    ("mkt_nb_cagr_fwd", "New build CAGR\n{cy}-{fy}", "0.0%", "Split", 12, "Optional", "in"),
    ("hfc_sales", "HFC sales {py}\n(EUR m)", "#,##0", "HFC", 12, "Actual sales in the base year", "in"),
    ("hfc_cagr_hist", "HFC sales CAGR\n{hy}-{py}", "0.0%", "HFC", 12, None, "in"),
    ("hfc_growth_cy", "HFC growth\n{cy} forecast", "0.0%", "HFC", 12, None, "in"),
    ("hfc_cagr_bp", "HFC BP CAGR\n{cy}-{fy}", "0.0%", "HFC", 12, "Business plan", "in"),
    ("ms_py_override", "Market share {py}\n(optional)", "0.0%", "HFC", 12, "Leave empty to use HFC sales / addressable market", "in"),
    ("ms_cy_override", "Market share {cy}\n(optional)", "0.0%", "HFC", 12, "Leave empty to grow sales and market by their current-year growth", "in"),
    ("profitability", "Profitability", "0", "Score", 12, "1-5 score: price level and margin potential of the market (5 = best)", "in"),
    ("competitive_intensity", "Competitive\nintensity", "0", "Score", 12, "1-5 score: 5 = benign competition, 1 = very fierce", "in"),
    ("channel_access", "Channel\naccess", "0", "Score", 12, "1-5 score: HFC access to distributors, installers, specifiers", "in"),
    ("product_fit", "Product\nfit", "0", "Score", 12, "1-5 score: fit of the HFC range with local norms and demand", "in"),
]
GROUPS = {
    "Country": ("Country / end market", INK_STRONG),
    "Market": ("Market (construction work done, Oxford Economics)", INK),
    "Split": ("Renovation / new build split (optional)", GREY_BAND),
    "HFC": ("HFC", RED),
    "Score": ("Qualitative scores (1-5, 5 = most favourable)", RED_SOFT),
}
COL = {c[0]: get_column_letter(i + 1) for i, c in enumerate(COLS)}
last = len(COLS)
band(rd, 1, last, "Raw data", "One row per country and end market. Blue cells are inputs; green cells are calculated from the Oxford Economics sheet.")
HEAD_ROW, FIRST = 5, 6
LAST_ROW = FIRST + N_ROWS - 1
start = 0
for i in range(1, last + 1):
    g = COLS[i - 1][3]
    nxt = COLS[i][3] if i < last else None
    if nxt != g:
        s_col, e_col = get_column_letter(start + 1), get_column_letter(i)
        if s_col != e_col:
            rd.merge_cells(f"{s_col}4:{e_col}4")
        c = rd[f"{s_col}4"]
        c.value = GROUPS[g][0]
        c.font = font(10, True, WHITE)
        c.alignment = CENTER
        for j in range(start + 1, i + 1):
            rd.cell(row=4, column=j).fill = fill(GROUPS[g][1])
        start = i
rd.row_dimensions[4].height = 20
for i, (key, head, fmt, grp, width, comment, kind) in enumerate(COLS, 1):
    rd.column_dimensions[get_column_letter(i)].width = width
    c = rd.cell(row=HEAD_ROW, column=i, value=year_formula(head, YEAR))
    c.font = font(9, True, INK if grp in ("Market", "Split") else RED)
    c.alignment = CENTER
    c.border = BOTTOM_RED
    if comment:
        c.comment = Comment(comment, "Template")
rd.row_dimensions[HEAD_ROW].height = 42

# ================================================================ Market levels (formulas)
ml = wb.create_sheet("Market levels")
ml.sheet_properties.tabColor = GREY_BAND
band(ml, 1, 16, "Market levels", "Construction work done per country and end market (US$ m, 2023 prices), calculated from the Oxford Economics sheet. Do not type here.")
ml.cell(row=3, column=1, value="Residential = Residential total - Multi family. Non-residential = Multi family + Commercial. Empty = year or series not in the export.").font = font(9, color=INK, italic=True)
ML_HEAD = 5
for j, h in enumerate(["Country", "End market", "Oxford Economics location"], 1):
    c = ml.cell(row=ML_HEAD, column=j, value=h)
    c.font = font(9, True, WHITE)
    c.fill = fill(INK)
    c.alignment = CENTER
ml.column_dimensions["A"].width = 17
ml.column_dimensions["B"].width = 15
ml.column_dimensions["C"].width = 20
YEAR_OFFSETS = list(range(-6, 6))   # current year - 6 ... current year + 5
for k, off in enumerate(YEAR_OFFSETS):
    col = 4 + k
    c = ml.cell(row=ML_HEAD, column=col, value=f"={YEAR}{off:+d}" if off else f"={YEAR}")
    c.font = font(9, True, WHITE)
    c.fill = fill(RED if off >= 0 else INK)
    c.alignment = CENTER
    c.number_format = "0"
    ml.column_dimensions[get_column_letter(col)].width = 11
YCOL = {off: get_column_letter(4 + k) for k, off in enumerate(YEAR_OFFSETS)}

def comp(ind_ref, year_cell, loc_cell):
    idx = f"IFERROR(MATCH({year_cell},{OE_HDR},0),MATCH({year_cell}&\"\",{OE_HDR},0))"
    return f"SUMIFS(INDEX({OE_DATA},0,{idx}),{OE_LOC},{loc_cell},{OE_IND},{ind_ref})"

for n in range(N_ROWS):
    mr, rr = ML_HEAD + 1 + n, FIRST + n
    calc_cell(ml.cell(row=mr, column=1, value=f"=IF('Raw data'!$A{rr}=\"\",\"\",'Raw data'!$A{rr})"), "@", 9)
    calc_cell(ml.cell(row=mr, column=2, value=f"=IF($A{mr}=\"\",\"\",'Raw data'!$C{rr})"), "@", 9)
    calc_cell(ml.cell(row=mr, column=3, value=f"=IF($A{mr}=\"\",\"\",IF('Raw data'!$D{rr}=\"\",$A{mr},'Raw data'!$D{rr}))"), "@", 9)
    for off in YEAR_OFFSETS:
        yc = f"{YCOL[off]}${ML_HEAD}"
        res = f"({comp(IND['oe.res_total'], yc, f'$C{mr}')}-{comp(IND['oe.res_mf'], yc, f'$C{mr}')})"
        non = f"({comp(IND['oe.res_mf'], yc, f'$C{mr}')}+{comp(IND['oe.nonres_com'], yc, f'$C{mr}')})"
        lvl = f'IF($B{mr}="Residential",{res},{non})'
        f = f'=IF($A{mr}="","",IFERROR(1/(1/{lvl}),""))'
        calc_cell(ml.cell(row=mr, column=4 + YEAR_OFFSETS.index(off), value=f), "#,##0", 9)
ml.freeze_panes = ml[f"D{ML_HEAD + 1}"]
ML = "'Market levels'!"

# ---------------------------------------------------------------- Raw data rows
sample = []
for name, code, loc, ms, hfc, scores in COUNTRIES:
    for s_i, seg in enumerate(("Residential", "Non-residential")):
        if seg == "Residential":
            lvl = oe_level(loc, IND_TOTAL, 2025) - oe_level(loc, IND_MF, 2025)
        else:
            lvl = oe_level(loc, IND_MF, 2025) + oe_level(loc, IND_COM, 2025)
        size = round(lvl * USD_EUR * ADDRESSABLE_RATIO / 5) * 5
        sales = round(size * ms[s_i] / 100, 1)
        bump = 0.5 if seg == "Non-residential" else 0.0
        sample.append({
            "country": name, "code": code, "segment": seg, "oe_location": loc,
            "market_size": size, "hfc_sales": sales,
            "hfc_cagr_hist": (hfc[0] + bump) / 100, "hfc_growth_cy": (hfc[1] + bump) / 100, "hfc_cagr_bp": (hfc[2] + bump) / 100,
            "profitability": scores[0], "competitive_intensity": scores[1],
            "channel_access": scores[2], "product_fit": scores[3],
        })

for n in range(N_ROWS):
    rr, mr = FIRST + n, ML_HEAD + 1 + n
    lv = lambda off: f"{ML}{YCOL[off]}{mr}"
    calc = {
        "oe_level_py": f'=IF($A{rr}="","",{lv(-1)})',
        "mkt_cagr_hist": f'=IF($A{rr}="","",IFERROR(({lv(-1)}/{lv(-6)})^(1/5)-1,""))',
        "mkt_growth_cy": f'=IF($A{rr}="","",IFERROR({lv(0)}/{lv(-1)}-1,""))',
        "mkt_cagr_fwd": f'=IF($A{rr}="","",IFERROR(({lv(5)}/{lv(0)})^(1/5)-1,""))',
    }
    for i, (key, head, fmt, grp, width, comment, kind) in enumerate(COLS, 1):
        c = rd.cell(row=rr, column=i)
        if kind == "calc":
            c.value = calc[key]
            calc_cell(c, fmt)
        else:
            input_cell(c, fmt)
    if n < len(sample):
        for key, v in sample[n].items():
            rd[f"{COL[key]}{rr}"].value = v
dv_seg = DataValidation(type="list", formula1='"Residential,Non-residential"', allow_blank=True)
dv_score = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                          error="Enter a whole score from 1 to 5", errorTitle="Score 1-5")
dv_pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                        error="Enter a percentage between 0% and 100%", errorTitle="Percentage")
for dv in (dv_seg, dv_score, dv_pct):
    rd.add_data_validation(dv)
dv_seg.add(f"{COL['segment']}{FIRST}:{COL['segment']}{LAST_ROW}")
dv_score.add(f"{COL['profitability']}{FIRST}:{COL['product_fit']}{LAST_ROW}")
for key in ("reno_share", "ms_py_override", "ms_cy_override"):
    dv_pct.add(f"{COL[key]}{FIRST}:{COL[key]}{LAST_ROW}")
rd.freeze_panes = rd[f"E{FIRST}"]
rd.cell(row=3, column=1, value=("Market growth comes from Oxford Economics (green). Addressable market, HFC figures and scores in the sample rows are ILLUSTRATIVE: "
                                "replace them with your own. Add countries in the empty rows.")).font = font(9, color=RED, italic=True)

# ================================================================ Dashboard data
dd = wb.create_sheet("Dashboard data")
dd.sheet_properties.tabColor = GREY_BAND
OUT_FIELDS = [
    ("country", "Country", "@"), ("code", "Code", "@"), ("segment", "End market", "@"),
    ("market_size", "Market size {py}", "#,##0"), ("hfc_sales", "HFC sales {py}", "#,##0"),
    ("mkt_cagr_hist", "Market CAGR {hy}-{py} (%)", "0.0"), ("mkt_growth_cy", "Market growth {cy}E (%)", "0.0"), ("mkt_cagr_fwd", "Market CAGR {cy}-{fy} (%)", "0.0"),
    ("hfc_cagr_hist", "HFC CAGR {hy}-{py} (%)", "0.0"), ("hfc_growth_cy", "HFC growth {cy}F (%)", "0.0"), ("hfc_cagr_bp", "HFC BP CAGR {cy}-{fy} (%)", "0.0"),
    ("ms_cy", "HFC share {cy} (%)", "0.0"), ("ms_py", "HFC share {py} (%)", "0.0"),
    ("profitability", "Profitability score", "0"), ("competitive_intensity", "Competitive intensity score", "0"),
    ("channel_access", "Channel access score", "0"), ("product_fit", "Product fit score", "0"),
    ("reno_share", "Renovation share (%)", "0"),
    ("mkt_reno_cagr_hist", "Reno CAGR {hy}-{py} (%)", "0.0"), ("mkt_reno_growth_cy", "Reno growth {cy}E (%)", "0.0"), ("mkt_reno_cagr_fwd", "Reno CAGR {cy}-{fy} (%)", "0.0"),
    ("mkt_nb_cagr_hist", "NB CAGR {hy}-{py} (%)", "0.0"), ("mkt_nb_growth_cy", "NB growth {cy}E (%)", "0.0"), ("mkt_nb_cagr_fwd", "NB CAGR {cy}-{fy} (%)", "0.0"),
]
band(dd, 1, len(OUT_FIELDS) + 1, "Dashboard data", "Calculated from Raw data and Settings. Do not type here: this is the table the dashboard reads.")
for i, (key, head, fmt) in enumerate(OUT_FIELDS, 1):
    dd.column_dimensions[get_column_letter(i)].width = 16 if i <= 3 else 13
    c = dd.cell(row=4, column=i, value=year_formula(head, YEAR))
    c.font = font(9, True, WHITE)
    c.fill = fill(RED if key.startswith(("hfc", "ms_")) else INK)
    c.alignment = CENTER
    k = dd.cell(row=5, column=i, value=key)
    k.font = font(8, color=MUTED)
    k.alignment = Alignment(horizontal="center")
dd.row_dimensions[4].height = 36
chk_i = len(OUT_FIELDS) + 1
chk_col = get_column_letter(chk_i)
dd.column_dimensions[chk_col].width = 34
c = dd.cell(row=4, column=chk_i, value="Check: renovation + new build vs. total, growth CY (pp)")
c.font = font(9, True, WHITE)
c.fill = fill(SALMON)
c.alignment = CENTER
dd.cell(row=5, column=chk_i, value="check").font = font(8, color=MUTED)

R = "'Raw data'!"
D0 = 6
for n in range(N_ROWS):
    r_out, rr = D0 + n, FIRST + n
    raw = lambda key: f"{R}{COL[key]}{rr}"
    formulas = {
        "country": f'=IF(TRIM({raw("country")})="","",TRIM({raw("country")}))',
        "code": f'=IF($A{r_out}="","",IF({raw("code")}="",UPPER(LEFT($A{r_out},3)),{raw("code")}))',
        "segment": f'=IF($A{r_out}="","",{raw("segment")})',
        "market_size": f'=IF(OR($A{r_out}="",{raw("market_size")}=""),"",{raw("market_size")})',
        "hfc_sales": f'=IF(OR($A{r_out}="",{raw("hfc_sales")}=""),"",{raw("hfc_sales")})',
        "ms_py": (f'=IF($A{r_out}="","",IF({raw("ms_py_override")}<>"",{raw("ms_py_override")}*100,'
                  f'IF(AND(N({raw("market_size")})>0,{raw("hfc_sales")}<>""),{raw("hfc_sales")}/{raw("market_size")}*100,"")))'),
        "ms_cy": (f'=IF($A{r_out}="","",IF({raw("ms_cy_override")}<>"",{raw("ms_cy_override")}*100,'
                  f'IF(AND(N({raw("market_size")})>0,{raw("hfc_sales")}<>""),'
                  f'{raw("hfc_sales")}*(1+N({raw("hfc_growth_cy")}))/({raw("market_size")}*(1+N({raw("mkt_growth_cy")})))*100,"")))'),
    }
    for key in ("profitability", "competitive_intensity", "channel_access", "product_fit"):
        formulas[key] = f'=IF($A{r_out}="","",IF(ISNUMBER({raw(key)}),{raw(key)},""))'
    for key in ("mkt_cagr_hist", "mkt_growth_cy", "mkt_cagr_fwd", "hfc_cagr_hist", "hfc_growth_cy", "hfc_cagr_bp", "reno_share",
                "mkt_reno_cagr_hist", "mkt_reno_growth_cy", "mkt_reno_cagr_fwd", "mkt_nb_cagr_hist", "mkt_nb_growth_cy", "mkt_nb_cagr_fwd"):
        formulas[key] = f'=IF($A{r_out}="","",IF(ISNUMBER({raw(key)}),{raw(key)}*100,""))'
    for i, (key, head, fmt) in enumerate(OUT_FIELDS, 1):
        calc_cell(dd.cell(row=r_out, column=i, value=formulas[key].replace("{r}", str(r_out))), fmt, 9)
    sh, g, rg, ng = raw("reno_share"), raw("mkt_growth_cy"), raw("mkt_reno_growth_cy"), raw("mkt_nb_growth_cy")
    chk = dd.cell(row=r_out, column=chk_i,
                  value=f'=IF($A{r_out}="","",IF(AND(ISNUMBER({sh}),ISNUMBER({g}),ISNUMBER({rg}),ISNUMBER({ng})),ROUND(({sh}*{rg}+(1-{sh})*{ng}-{g})*100,1),""))')
    chk.font = font(9)
    chk.number_format = "+0.0;-0.0;0.0"
dd.conditional_formatting.add(
    f"{chk_col}{D0}:{chk_col}{D0 + N_ROWS - 1}",
    FormulaRule(formula=[f'AND(ISNUMBER({chk_col}{D0}),ABS({chk_col}{D0})>0.5)'], fill=fill(BLUSH), font=Font(name=FONT, color=RED, bold=True)))
dd.freeze_panes = dd[f"D{D0}"]
dd.cell(row=3, column=1, value="Percentages are plain numbers (3.2 = 3.2%). Check column: more than ±0.5 pp means the renovation / new build split does not add up to the total.").font = font(9, color=INK, italic=True)

# ================================================================ Read me
rm = wb.create_sheet("Read me", 0)
rm.sheet_properties.tabColor = RED
band(rm, 1, 8, "Market intelligence Dashboard - data template", "Update the blue cells, save, and upload the file in the dashboard")
rm.column_dimensions["A"].width = 4
rm.column_dimensions["B"].width = 30
for col in "CDEFGH":
    rm.column_dimensions[col].width = 16
lines = [
    ("How to update the dashboard", None),
    ("1", "Oxford Economics: paste the latest download (level values) at cell A6, header row included. Market growth for every country updates automatically."),
    ("2", "Raw data: one row per country and end market. Update the blue cells: addressable market, HFC sales and growth, and the four 1-5 scores."),
    ("3", "Settings: check the company name, the current year and the Oxford Economics series names. Weights and zone lines are set in the dashboard."),
    ("4", "Dashboard data: nothing to type. Check that every country appears and that the check column shows no red cells."),
    ("5", "Save the file (.xlsx), open the dashboard, go to 'Update the figures' and click 'Upload Excel'."),
    ("How the market figures are built", None),
    ("End markets", "Residential = Oxford Economics 'Residential - Total' minus 'Multi Family'. Non-residential = 'Multi Family' plus 'Non-Residential Building - Commercial'."),
    ("Market levels", "Shows the resulting construction work done per country, end market and year (US$ m, 2023 prices)."),
    ("Growth rates", "CAGR last 5 years = (base year / base year - 5)^(1/5) - 1. Growth this year = current year / base year - 1. CAGR next 5 years = (current year + 5 / current year)^(1/5) - 1."),
    ("History", "The 5-year history needs the export to start 6 years before the current year (e.g. 2020 for 2026). With a shorter export it stays empty."),
    ("Addressable market", "Your estimate of the HFC addressable market in EUR m (not total construction output). Used for market size and market share."),
    ("Renovation / new build", "Optional inputs. When empty, the dashboard uses total market growth in the matrix and hides the split."),
    ("Qualitative scores", "Profitability, competitive intensity, channel access and product fit: type a whole score from 1 to 5 (5 = most favourable for HFC)."),
    ("Conventions", None),
    ("Blue cells", "Inputs: type or paste your figures here."),
    ("Green cells", "Formulas: do not overwrite."),
    ("Percentages", "Type them as percentages (3.2%) or decimals (0.032). Growth rates can be negative."),
    ("Rows", f"Room for {N_ROWS} rows in Raw data (for example {N_ROWS // 2} countries x 2 end markets). Empty rows are ignored."),
]
row = 4
for a, b in lines:
    if b is None:
        row += 1
        section(rm, row, a, col=2)
        rm.cell(row=row, column=2).border = BOTTOM_RED
        row += 1
        continue
    ca = rm.cell(row=row, column=2, value=a)
    ca.font = font(10, True, RED if a.isdigit() else INK_STRONG)
    ca.alignment = Alignment(horizontal="left", vertical="center")
    cb = rm.cell(row=row, column=3, value=b)
    cb.font = font(10, color=INK)
    rm.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
    cb.alignment = WRAP
    rm.row_dimensions[row].height = 32
    if a == "Blue cells":
        ca.font = font(10, True, INPUT_BLUE)
        ca.fill = fill(TINT)
    if a == "Green cells":
        ca.font = font(10, True, LINK_GREEN)
    row += 1
rm.cell(row=row + 1, column=2, value=("Market growth is real Oxford Economics data. Addressable market, HFC figures and qualitative scores in the sample rows "
                                      "are illustrative, not HFC data.")).font = font(9, color=RED, italic=True)
for ws in (rm, st):
    ws.sheet_view.showGridLines = False
for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 90
wb.active = 0

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
